# pii_contracte - detectarea și redactarea datelor de identificare din contracte.
# Copyright (C) 2026  CatalinFusea
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published
# by the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.
# SPDX-License-Identifier: AGPL-3.0-or-later

"""Generare de date și documente SINTETICE pentru teste. Nicio dată reală."""
from __future__ import annotations

import copy
import os
import pathlib

import docx
import pymupdf
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",          # Linux (Debian/Ubuntu)
    os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts", "arial.ttf"),   # Windows
    "/System/Library/Fonts/Supplemental/Arial.ttf",              # macOS
]
FONT = next((f for f in _FONT_CANDIDATES if os.path.isfile(f)), None)
if FONT is None:
    raise RuntimeError("Niciun font TTF cu diacritice găsit pentru teste: " + ", ".join(_FONT_CANDIDATES))


def gen_cnp(s: int, yy: int, mm: int, dd: int, county: int, nnn: int) -> str:
    body = f"{s}{yy:02d}{mm:02d}{dd:02d}{county:02d}{nnn:03d}"
    t = sum(int(a) * int(b) for a, b in zip(body, "279146358279")) % 11
    return body + str(1 if t == 10 else t)


def gen_iban_ro(bank: str, account: str) -> str:
    assert len(bank) == 4 and len(account) == 16
    bban = bank + account
    num = int("".join(str(int(c, 36)) for c in bban + "RO00"))
    return f"RO{98 - num % 97:02d}{bban}"


CNP_A = gen_cnp(1, 80, 1, 15, 40, 123)     # București
CNP_B = gen_cnp(2, 92, 7, 3, 23, 45)       # Ilfov
CNP_OLD = gen_cnp(1, 75, 3, 3, 12, 7)      # apare doar în textul șters (track changes)
IBAN_A = gen_iban_ro("BTRL", "0000123456789012")
NAME_A = "POPESCU ION-MARIAN"
NAME_B = "Ionescu Maria"
EMAIL_A = "ion.popescu@example.com"
EMAIL_B = "maria.ionescu@example.org"
PHONE_A = "0722 123 456"
CONTRACT_NO = "1234567890123"  # nr. de contract cu 13 cifre, NU e CNP valid


def preamble() -> str:
    return (f"Între: {NAME_A}, domiciliat în Mun. București, Str. Florilor nr. 5, bl. A2, sc. 1, ap. 14, "
            f"Sector 3, CNP {CNP_A}, identificat cu CI seria RX nr. 123456, tel. {PHONE_A}, "
            f"în calitate de vânzător, și")


def preamble_b() -> str:
    return (f"Doamna {NAME_B}, cu domiciliul în Jud. Ilfov, Com. Otopeni, Str. Mare nr. 2, "
            f"CNP {CNP_B}, în calitate de cumpărător.")


def make_docx(path: str) -> None:
    d = docx.Document()
    d.core_properties.author = "Popescu Ion"
    d.core_properties.last_modified_by = "Popescu Ion"
    d.add_heading("CONTRACT DE VÂNZARE-CUMPĂRARE", 0)
    d.add_paragraph(f"Contract nr. {CONTRACT_NO} din 01.02.2026")
    # CNP-ul spart în trei run-uri (cum face Word)
    p = d.add_paragraph()
    pre = preamble()
    i = pre.index(CNP_A)
    p.add_run(pre[:i + 4])
    p.add_run(CNP_A[4:9]).bold = True
    p.add_run(pre[i + 9:])
    d.add_paragraph(preamble_b())
    # track changes: CNP vechi șters, inserat text nou
    p2 = d.add_paragraph("Rectificare: ")
    dele = OxmlElement("w:del")
    dele.set(qn("w:id"), "1"); dele.set(qn("w:author"), "Popescu Ion"); dele.set(qn("w:date"), "2026-01-01T00:00:00Z")
    r = OxmlElement("w:r"); t = OxmlElement("w:delText"); t.text = f"CNP {CNP_OLD}"
    t.set(qn("xml:space"), "preserve"); r.append(t); dele.append(r)
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), "2"); ins.set(qn("w:author"), "Popescu Ion"); ins.set(qn("w:date"), "2026-01-01T00:00:00Z")
    r2 = OxmlElement("w:r"); t2 = OxmlElement("w:t"); t2.text = "fără modificări."; r2.append(t2); ins.append(r2)
    p2._p.append(dele); p2._p.append(ins)
    # tabel cu IBAN
    tb = d.add_table(rows=1, cols=2)
    tb.cell(0, 0).text = "Cont"
    tb.cell(0, 1).text = f"IBAN {IBAN_A}"
    # hyperlink mailto
    p3 = d.add_paragraph("Contact: ")
    rid = d.part.relate_to(f"mailto:{EMAIL_B}", RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    hr = OxmlElement("w:r"); ht = OxmlElement("w:t"); ht.text = EMAIL_B; hr.append(ht); h.append(hr)
    p3._p.append(h)
    # comentariu cu telefon
    cp = d.add_paragraph("Prețul este de 1.000 lei.")
    d.add_comment(cp.runs, text="Sunați-l pe 0733 444 555", author="Ionescu Maria", initials="IM")
    # antet cu email
    d.sections[0].header.paragraphs[0].text = f"Redactat de {EMAIL_A}"
    # semnături: numele fără ancoră -> trebuie găsit prin propagare
    d.add_paragraph("Vânzător,")
    d.add_paragraph(NAME_A)
    d.add_paragraph("Clauză fără date personale. Contract de vânzare.")
    d.save(path)


def make_pdf(path: str, with_scanned_page: bool = True) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    text = (f"CONTRACT DE VÂNZARE-CUMPĂRARE\nContract nr. {CONTRACT_NO}\n"
            f"{preamble()}\n{preamble_b()}\nIBAN: {IBAN_A}\nEmail: {EMAIL_B}\n\n"
            f"Clauză fără date personale. Contract de vânzare.\n\nVânzător,\n{NAME_A}")
    rect = pymupdf.Rect(50, 50, 545, 800)
    rc = page.insert_textbox(rect, text, fontname="dejavu", fontfile=FONT, fontsize=10)
    assert rc >= 0, "textul nu a încăput"
    page.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(50, 780, 300, 795),
                      "uri": f"mailto:{EMAIL_B}"})
    if with_scanned_page:
        p2 = doc.new_page()
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 100), False)
        pix.clear_with(200)
        p2.insert_image(pymupdf.Rect(50, 50, 450, 250), pixmap=pix)
    doc.set_metadata({"author": "Popescu Ion", "title": "Contract Popescu"})
    doc.save(path)
    doc.close()


def make_encrypted_pdf(path: str) -> None:
    doc = pymupdf.open()
    doc.new_page().insert_text((50, 50), "secret")
    doc.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="x", owner_pw="y")
    doc.close()


def make_xxe_docx(src: str, dst: str, target: str) -> None:
    """Copie a unui docx în care document.xml declară o entitate externă spre `target`."""
    import zipfile
    uri = pathlib.Path(target).resolve().as_uri().encode()
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w") as zo:
        for info in zi.infolist():
            data = zi.read(info.filename)
            if info.filename == "word/document.xml":
                data = data.replace(b"?>", b'?><!DOCTYPE x [<!ENTITY xxe SYSTEM "' + uri + b'">]>', 1)
                data = data.replace(b"Clauz", b"&xxe;Clauz", 1)
            zo.writestr(info, data)
    _ = copy  # păstrat pentru extensii viitoare
