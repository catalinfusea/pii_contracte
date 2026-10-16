"""2.3: permis, CF/cadastral, adresa imobilului (bifate implicit); loc naștere, emitent/date acte (nebifate).
Text-model după un contract de închiriere real; toate valorile sunt SINTETICE."""
import docx
import pytest

import fixtures as fx
from pii_core import detectors as det
from pii_core import engine

CNP_C = fx.gen_cnp(2, 85, 4, 12, 8, 31)
P1 = (f"1. Subsemnata, VASILESCU ELENA, cetățean român, născută la data de 12.04.1985 în Mun. Sibiu, "
      f"jud. Sibiu, domiciliată în Mun. Cluj-Napoca, Str. Teilor nr. 3, ap. 7, identificată cu C.I. seria "
      f"KX nr. 654321, eliberată de SPCLEP Cluj-Napoca la data de 02.03.2016, valabilă până la 12.04.2026, "
      f"CNP {CNP_C}, telefon 0744 555 666, în calitate de PROPRIETAR (LOCATOR).")
P2 = ("2. Subsemnatul, DUMITRU ANDREI, cetățean român, necăsătorit, domiciliat în Mun. Iași, Str. Lungă nr. 9, "
      "identificat prin pașaport tip P, nr. 051234567, eliberat de autoritățile române la data de 05.06.2020, "
      "posesor al permisului de conducere categoria B cu numărul Z00998877 eliberat de DRPCIV Iași, "
      "în calitate de CHIRIAȘ (LOCATAR).")
A1 = ("Art. 1. Obiectul contractului constă în închirierea imobilului situat în București, Strada Merilor "
      "nr. 21, bloc C3, etaj 4, ap. 18, Sector 2, înregistrat în Cartea Funciară sub nr. 98765, compus din "
      "2 camere, bucătărie și baie.")
A2 = "Art. 2. Plata se face la data de 10 a lunii. Contractul a fost semnat la data de 01.09.2023."


def kv(text):
    return {(m.kind, m.value) for m in det.detect(text)}


def test_property_and_licence():
    got = kv(A1) | kv(P2)
    assert ("Adresă imobil", "București, Strada Merilor nr. 21, bloc C3, etaj 4, ap. 18, Sector 2") in got
    assert ("Carte funciară / cadastral", "98765") in got
    assert ("Număr permis conducere", "Z00998877") in got
    assert ("Emitent act", "DRPCIV Iași") in got


def test_indirect():
    got = kv(P1) | kv(P2)
    assert ("Loc naștere", "Mun. Sibiu, jud. Sibiu") in got
    assert ("Emitent act", "SPCLEP Cluj-Napoca") in got
    assert ("Dată act", "02.03.2016") in got and ("Dată act", "12.04.2026") in got
    assert ("Dată act", "05.06.2020") in got
    assert ("Data nașterii", "12.04.1985") in got               # neschimbat
    assert ("Adresă", "Mun. Cluj-Napoca, Str. Teilor nr. 3, ap. 7") in got


@pytest.mark.parametrize("text", [
    A2,
    "Art. 3. Contractul se încheie pe 24 de luni, începând cu 01.09.2023 până la 31.08.2025.",
    "Chiria se plătește în contul menționat la Capitolul I. Garanția este de 900 EUR.",
    "Certificatul energetic este anexat. Prețul se actualizează anual.",
])
def test_no_false_positives(text):
    assert kv(text) == set()


def test_cadastral_variants():
    assert ("Carte funciară / cadastral", "123456") in kv("imobil cu număr cadastral 123456, CF nr. 123456")
    assert ("Carte funciară / cadastral", "201811-C1-U5") in kv("înscris în CF nr. 201811-C1-U5 a localității")


def test_defaults_in_scan(tmp_path):
    p = tmp_path / "inchiriere.docx"
    d = docx.Document()
    for t in (P1, P2, A1, A2, "LOCATOR,", "VASILESCU ELENA", "LOCATAR,", "DUMITRU ANDREI"):
        d.add_paragraph(t)
    d.save(str(p))
    res = engine.scan_file(str(p))
    sel = {f.kind for f in res.findings if f.selected}
    unsel = {f.kind for f in res.findings if not f.selected}
    assert {"Adresă imobil", "Carte funciară / cadastral", "Număr permis conducere"} <= sel
    assert unsel == {"Loc naștere", "Emitent act", "Dată act"}
    # semnăturile prinse prin propagare
    assert sum(f.value == "VASILESCU ELENA" for f in res.findings) == 2
    out, problems = engine.redact_file(res, str(tmp_path / "o"))
    assert problems == []
    text = "\n".join(x.text for x in docx.Document(out).paragraphs)
    for v in ("Z00998877", "98765", "Merilor", "VASILESCU", "DUMITRU", CNP_C):
        assert v not in text, v
    for v in ("Mun. Sibiu", "SPCLEP", "02.03.2016", "DRPCIV", "01.09.2023"):   # nebifate -> rămân
        assert v in text, v
