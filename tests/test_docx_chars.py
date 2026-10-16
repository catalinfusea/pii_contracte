"""Cratime stocate de Word ca elemente (w:noBreakHyphen, w:softHyphen), nu ca text."""
import zipfile

import docx
from docx.oxml import OxmlElement

import fixtures as fx
from pii_core import engine


def _doc(path, sep_tag, tail=""):
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("Între: POPESCU ION")
    r = p.add_run()
    r._r.append(OxmlElement(sep_tag))
    p.add_run(f"MARIAN, domiciliat în Mun. București, Str. Florilor nr. 5, CNP {fx.CNP_A}")
    d.add_paragraph(f"Semnătura: POPESCU ION-MARIAN{tail}")      # cratimă normală, fără ancoră
    d.add_paragraph("Termen: 30 de zile")
    q = d.add_paragraph()
    q.add_run("Text non")
    q.add_run()._r.append(OxmlElement("w:noBreakHyphen"))
    q.add_run("personal rămâne.")
    d.save(path)


def test_nobreak_hyphen(tmp_path):
    src = str(tmp_path / "a.docx")
    _doc(src, "w:noBreakHyphen")
    res = engine.scan_file(src)
    names = [f for f in res.findings if f.kind == "Nume"]
    assert names and all(f.value == "POPESCU ION-MARIAN" for f in names)
    assert len(names) == 2                       # propagarea găsește și semnătura
    out, problems = engine.redact_file(res, str(tmp_path / "o"))
    assert problems == []
    paras = [p.text for p in docx.Document(out).paragraphs]
    assert paras[0].startswith("Între: [REDACTAT:NUME], domiciliat")   # fără cratimă rătăcită
    assert paras[1] == "Semnătura: [REDACTAT:NUME]"
    xml = zipfile.ZipFile(out).read("word/document.xml").decode()
    assert "POPESCU" not in xml and "MARIAN" not in xml
    assert xml.count("noBreakHyphen") == 1        # cea din textul ne-personal rămâne


def test_soft_hyphen_inside_name(tmp_path):
    src = str(tmp_path / "b.docx")
    _doc(src, "w:softHyphen")
    res = engine.scan_file(src)
    assert ("Nume", "POPESCU IONMARIAN") in {(f.kind, f.value) for f in res.findings}  # textul vizibil
    out, problems = engine.redact_file(res, str(tmp_path / "o"))
    assert problems == []
    xml = zipfile.ZipFile(out).read("word/document.xml").decode()
    assert "softHyphen" not in xml
    first = docx.Document(out).paragraphs[0].text
    assert first.startswith("Între: [REDACTAT:NUME], domiciliat") and "POPESCU" not in first
    # limită cunoscută: altă grafie a numelui (cu cratimă) nu e legată de propagare -> nu e redactată
