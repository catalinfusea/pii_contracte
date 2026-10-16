# pii_contracte - detectarea și redactarea datelor de identificare din contracte.
# Copyright (C) 2026  CatalinFusea
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published
# by the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.
# SPDX-License-Identifier: AGPL-3.0-or-later

"""Export după redactare (DOCX -> PDF, PDF -> DOCX) cu un convertor FALS; Word real: checklist."""
import os
import queue
import threading

import docx
import pymupdf
import pytest

import fixtures as fx
from pii_core import engine
from pii_gui import jobs


class TextConverter:
    """Simulează Word: scrie în formatul cerut textul vizibil al copiei redactate (+ opțional o scurgere)."""
    def __init__(self, leak=None):
        self.leak, self.calls = leak, []

    def __enter__(self):
        return self

    def __exit__(self, *a):
        self.calls.append("exit")

    def export(self, src, fmt, dst):
        self.calls.append((os.path.splitext(src)[1], fmt))
        if src.endswith(".docx"):
            text = "\n".join(p.text for p in docx.Document(src).paragraphs)
        else:
            d = pymupdf.open(src)
            text = "".join(p.get_text() for p in d)
            d.close()
        if self.leak:
            text += "\n" + self.leak
        if fmt == "pdf":
            d = pymupdf.open()
            d.new_page().insert_textbox(pymupdf.Rect(40, 40, 560, 800), text, fontname="f0",
                                        fontfile=fx.FONT, fontsize=8)
            d.save(dst)
            d.close()
        else:
            d = docx.Document()
            for line in text.splitlines():
                d.add_paragraph(line)
            d.save(dst)
        return dst


@pytest.fixture
def results(tmp_path):
    a, b = tmp_path / "a.docx", tmp_path / "a.pdf"           # același nume, formate diferite
    fx.make_docx(str(a))
    fx.make_pdf(str(b))
    return [engine.scan_file(str(a)), engine.scan_file(str(b))]


def run(results, out, fmt, conv):
    q = queue.Queue()
    jobs.redact_job(results, str(out), q, threading.Event(), out_format=fmt, converter_factory=lambda: conv)
    msgs = []
    while not q.empty():
        msgs.append(q.get())
    return msgs


def test_export_pdf(results, tmp_path):
    conv = TextConverter()
    msgs = run(results, tmp_path / "o", "pdf", conv)
    names = sorted(os.listdir(tmp_path / "o"))
    # DOCX redactat + PDF exportat; PDF-ul redactat direct nu e suprascris de export
    assert "a_REDACTAT.docx" in names and "a_REDACTAT.pdf" in names and "a_REDACTAT_2.pdf" in names
    assert conv.calls == [(".docx", "pdf"), "exit"]           # PDF-ul nu se mai exportă
    done = msgs[-1][1]
    assert done["done"] == 2 and done["problems"] == 0
    assert len([m for m in msgs if m[0] == "redacted"]) == 3


def test_export_docx(results, tmp_path):
    conv = TextConverter()
    run(results, tmp_path / "o", "docx", conv)
    names = os.listdir(tmp_path / "o")
    assert "a_REDACTAT.pdf" in names and "a_REDACTAT_2.docx" in names
    assert conv.calls == [(".pdf", "docx"), "exit"]


def test_original_no_word(results, tmp_path):
    def boom():
        raise AssertionError("Word nu trebuie pornit")
    msgs = run(results, tmp_path / "o", "original", boom)
    assert msgs[-1][1]["done"] == 2


def test_leak_in_export_is_reported(results, tmp_path):
    msgs = run(results[:1], tmp_path / "o", "pdf", TextConverter(leak=f"CNP {fx.CNP_A}"))
    exported = [m for m in msgs if m[0] == "redacted" and m[2].endswith(".pdf")]
    assert exported and any("CNP" in p for p in exported[0][3])
    assert msgs[-1][1]["problems"] == 1


def test_word_unavailable_keeps_native_copy(results, tmp_path):
    from pii_core import doc_convert

    class NoWord:
        def __enter__(self):
            raise doc_convert.ConversionError("Microsoft Word nu poate fi pornit (test).")
    msgs = run(results, tmp_path / "o", "pdf", NoWord())
    errs = [m for m in msgs if m[0] == "file_error"]
    assert len(errs) == 1 and "rămas în formatul inițial" in errs[0][2]   # Word încercat o singură dată
    assert "a_REDACTAT.docx" in os.listdir(tmp_path / "o")
